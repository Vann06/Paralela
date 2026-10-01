from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def fmt_length(value):
    if value is None:
        return None
    return round(value.inches, 3)


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(path)
    print(f"FILE: {path}")
    print(f"PARAGRAPHS: {len(doc.paragraphs)} TABLES: {len(doc.tables)} SHAPES: {len(doc.inline_shapes)}")

    for i, section in enumerate(doc.sections, 1):
        print(
            f"SECTION {i}: size={fmt_length(section.page_width)}x{fmt_length(section.page_height)} "
            f"margins=({fmt_length(section.top_margin)}, {fmt_length(section.right_margin)}, "
            f"{fmt_length(section.bottom_margin)}, {fmt_length(section.left_margin)}) "
            f"header={fmt_length(section.header_distance)} footer={fmt_length(section.footer_distance)}"
        )

    print("\nSTYLES USED")
    used = {}
    for paragraph in doc.paragraphs:
        used[paragraph.style.name] = used.get(paragraph.style.name, 0) + 1
    for name, count in sorted(used.items()):
        style = doc.styles[name]
        font = style.font
        print(
            f"{name!r}: count={count}, font={font.name!r}, size={font.size.pt if font.size else None}, "
            f"bold={font.bold}, italic={font.italic}, color={font.color.rgb if font.color and font.color.rgb else None}"
        )

    print("\nPARAGRAPHS")
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.replace("\t", "[TAB]").replace("\n", "[NL]")
        ppr = paragraph._p.pPr
        sect = bool(ppr is not None and ppr.find(qn("w:sectPr")) is not None)
        print(
            f"P{i:03d} style={paragraph.style.name!r} align={paragraph.alignment} "
            f"break_before={paragraph.paragraph_format.page_break_before} sect={sect} text={text!r}"
        )
        for j, run in enumerate(paragraph.runs):
            if run.text or run._r.xpath('.//w:drawing'):
                print(
                    f"  R{j:02d} text={run.text!r} font={run.font.name!r} "
                    f"size={run.font.size.pt if run.font.size else None} bold={run.bold} "
                    f"italic={run.italic} underline={run.underline} color={run.font.color.rgb if run.font.color and run.font.color.rgb else None}"
                )

    print("\nTABLES")
    for ti, table in enumerate(doc.tables):
        print(f"TABLE {ti}: style={table.style.name if table.style else None}, rows={len(table.rows)}, cols={len(table.columns)}")
        for ri, row in enumerate(table.rows):
            print(f"  ROW {ri}: " + " || ".join(cell.text.replace("\n", " / ") for cell in row.cells))

    print("\nHEADERS/FOOTERS")
    for si, section in enumerate(doc.sections, 1):
        print(f"SECTION {si} HEADER: {[p.text for p in section.header.paragraphs]}")
        print(f"SECTION {si} FOOTER: {[p.text for p in section.footer.paragraphs]}")

    print("\nBODY ELEMENTS")
    for i, element in enumerate(doc._element.body):
        text = "".join(element.xpath('.//*[local-name()="t"]/text()'))
        breaks = [node.get(qn("w:type")) for node in element.xpath('.//*[local-name()="br"]')]
        page_break_before = bool(
            element.xpath('./*[local-name()="pPr"]/*[local-name()="pageBreakBefore"]')
        )
        print(
            i,
            element.tag.split("}")[-1],
            repr(text[:120]),
            "breaks=",
            breaks,
            "page_break_before=",
            page_break_before,
        )


if __name__ == "__main__":
    main()
